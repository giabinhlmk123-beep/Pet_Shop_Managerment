import customtkinter as ctk
from tkinter import ttk, messagebox

import os

import qrcode
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as ReportLabImage
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from database.database import connect_database
from models.session import Session


class CashForm:

    def __init__(self, parent):

        self.parent = parent

        self.cart = []

        self.create_interface()

        self.load_customers()
        self.load_products()
        self.load_invoice_history()

    # TẠO GIAO DIỆN

    def create_interface(self):

        # ==========================================
        # MÀU SẮC GIAO DIỆN GIỐNG DASHBOARD
        # ==========================================

        bg_color = "#F1F2F7"
        card_color = "#FFFFFF"
        text_color = "#202044"
        purple_color = "#6500F5"
        hover_color = "#4D00C4"
        border_color = "#E5E7F0"
        muted_color = "#85859B"
        primary_color = "#6500F5"

        self.parent.configure(fg_color=bg_color)

        # ==========================================
        # TIÊU ĐỀ
        # ==========================================

        header_frame = ctk.CTkFrame(
            self.parent,
            fg_color=bg_color
        )
        header_frame.pack(
            fill="x",
            padx=25,
            pady=(16, 8)
        )

        ctk.CTkLabel(
            header_frame,
            text="Quản lý bán hàng",
            font=("Arial", 27, "bold"),
            text_color=text_color
        ).pack(side="left")

        ctk.CTkLabel(
            header_frame,
            text="PET SHOP  /  SALES",
            font=("Arial", 11),
            text_color=muted_color
        ).pack(side="right", padx=5)

        # ==========================================
        # THÔNG TIN HÓA ĐƠN
        # ==========================================

        info_frame = ctk.CTkFrame(
            self.parent,
            fg_color=card_color,
            corner_radius=12,
            border_width=1,
            border_color=border_color
        )
        info_frame.pack(
            fill="x",
            padx=25,
            pady=6
        )

        ctk.CTkLabel(
            info_frame,
            text="Thông tin hóa đơn",
            font=("Arial", 16, "bold"),
            text_color=text_color
        ).grid(
            row=0,
            column=0,
            columnspan=4,
            padx=16,
            pady=(12, 4),
            sticky="w"
        )

        info_frame.grid_columnconfigure(1, weight=1)
        info_frame.grid_columnconfigure(3, weight=1)

        ctk.CTkLabel(
            info_frame,
            text="Khách hàng:",
            text_color=text_color
        ).grid(
            row=1,
            column=0,
            padx=16,
            pady=10,
            sticky="w"
        )

        self.cbo_customer = ctk.CTkComboBox(
            info_frame,
            width=300,
            height=34,
            values=[""],
            corner_radius=8,
            fg_color="#FAFAFC",
            border_color=border_color,
            text_color=text_color,
            button_color=primary_color,
            button_hover_color=hover_color,
            dropdown_fg_color=card_color,
            dropdown_text_color=text_color,
            command=self.customer_selected
        )
        self.cbo_customer.grid(
            row=1,
            column=1,
            padx=10,
            pady=10,
            sticky="ew"
        )

        ctk.CTkLabel(
            info_frame,
            text="Nhân viên:",
            text_color=text_color
        ).grid(
            row=1,
            column=2,
            padx=16,
            pady=10,
            sticky="w"
        )

        self.txt_cashier = ctk.CTkEntry(
            info_frame,
            width=250,
            height=34,
            corner_radius=8,
            fg_color="#FAFAFC",
            text_color="#000000",
            border_color=border_color
        )
        self.txt_cashier.grid(
            row=1,
            column=3,
            padx=10,
            pady=10,
            sticky="ew"
        )

        self.txt_cashier.insert(0, Session.full_name)
        self.txt_cashier.configure(state="disabled")

        # ==========================================
        # CHỌN SẢN PHẨM
        # ==========================================

        product_frame = ctk.CTkFrame(
            self.parent,
            fg_color=card_color,
            corner_radius=12,
            border_width=1,
            border_color=border_color
        )
        product_frame.pack(
            fill="x",
            padx=25,
            pady=6
        )

        ctk.CTkLabel(
            product_frame,
            text="Thông tin sản phẩm",
            font=("Arial", 16, "bold"),
            text_color=text_color
        ).grid(
            row=0,
            column=0,
            columnspan=9,
            padx=16,
            pady=(12, 4),
            sticky="w"
        )

        product_frame.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(
            product_frame,
            text="Sản phẩm:",
            text_color=text_color
        ).grid(
            row=1,
            column=0,
            padx=12,
            pady=10,
            sticky="w"
        )

        self.cbo_product = ctk.CTkComboBox(
            product_frame,
            width=240,
            height=34,
            values=[""],
            corner_radius=8,
            border_color=border_color,
            fg_color="#FAFAFC",
            text_color=text_color,
            button_color=primary_color,
            button_hover_color=hover_color,
            dropdown_fg_color=card_color,
            dropdown_text_color=text_color,
            command=self.product_selected
        )
        self.cbo_product.grid(
            row=1,
            column=1,
            padx=8,
            pady=10,
            sticky="ew"
        )

        ctk.CTkLabel(
            product_frame,
            text="Số lượng:",
            text_color=text_color
        ).grid(
            row=1,
            column=2,
            padx=8,
            pady=10
        )

        self.txt_quantity = ctk.CTkEntry(
            product_frame,
            width=75,
            height=34,
            corner_radius=8,
            fg_color="#FAFAFC",
            text_color="#000000",
            border_color=border_color
        )
        self.txt_quantity.grid(
            row=1,
            column=3,
            padx=8,
            pady=10
        )
        self.txt_quantity.insert(0, "1")

        ctk.CTkLabel(
            product_frame,
            text="Đơn giá:",
            text_color=text_color
        ).grid(
            row=1,
            column=4,
            padx=8,
            pady=10
        )

        self.txt_price = ctk.CTkEntry(
            product_frame,
            width=125,
            height=34,
            corner_radius=8,
            fg_color="#FAFAFC",
            text_color="#000000",
            border_color=border_color
        )
        self.txt_price.grid(
            row=1,
            column=5,
            padx=8,
            pady=10
        )
        self.txt_price.configure(state="disabled")

        ctk.CTkLabel(
            product_frame,
            text="Tồn kho:",
            text_color=text_color
        ).grid(
            row=1,
            column=6,
            padx=8,
            pady=10
        )

        self.lbl_stock = ctk.CTkLabel(
            product_frame,
            text="0",
            font=("Arial", 14, "bold"),
            text_color=purple_color
        )
        self.lbl_stock.grid(
            row=1,
            column=7,
            padx=8,
            pady=10
        )

        ctk.CTkButton(
            product_frame,
            text="＋ Thêm vào giỏ",
            width=135,
            height=34,
            corner_radius=8,
            fg_color=purple_color,
            hover_color=hover_color,
            command=self.add_to_cart
        ).grid(
            row=1,
            column=8,
            padx=12,
            pady=10
        )

        # ==========================================
        # KHUNG GIỎ HÀNG
        # ==========================================

        cart_frame = ctk.CTkFrame(
            self.parent,
            fg_color=card_color,
            corner_radius=12,
            border_width=1,
            border_color=border_color
        )
        cart_frame.pack(
            fill="both",
            expand=True,
            padx=25,
            pady=6
        )

        cart_header = ctk.CTkFrame(
            cart_frame,
            fg_color=card_color
        )
        cart_header.pack(
            fill="x",
            padx=15,
            pady=(12, 5)
        )

        ctk.CTkLabel(
            cart_header,
            text="Giỏ hàng",
            font=("Arial", 17, "bold"),
            text_color=text_color
        ).pack(side="left")

        # ==========================================
        # STYLE BẢNG
        # ==========================================

        tree_style = ttk.Style()

        tree_style.configure(
            "Cash.Treeview",
            background="#FFFFFF",
            foreground=text_color,
            fieldbackground="#FFFFFF",
            rowheight=30,
            font=("Arial", 10)
        )

        tree_style.configure(
            "Cash.Treeview.Heading",
            background="#F0EBFF",
            foreground=text_color,
            font=("Arial", 10, "bold"),
            padding=(6, 8)
        )

        tree_style.map(
            "Cash.Treeview",
            background=[("selected", "#E4D6FF")],
            foreground=[("selected", text_color)]
        )

        # ==========================================
        # BẢNG GIỎ HÀNG
        # ==========================================

        table_frame = ctk.CTkFrame(
            cart_frame,
            fg_color=card_color
        )
        table_frame.pack(
            fill="both",
            expand=True,
            padx=12,
            pady=5
        )

        columns = (
            "pcode",
            "pname",
            "qty",
            "price",
            "total"
        )

        self.tree_cart = ttk.Treeview(
            table_frame,
            columns=columns,
            show="headings",
            style="Cash.Treeview"
        )

        self.tree_cart.heading("pcode", text="Mã SP")
        self.tree_cart.heading("pname", text="Tên sản phẩm")
        self.tree_cart.heading("qty", text="Số lượng")
        self.tree_cart.heading("price", text="Đơn giá")
        self.tree_cart.heading("total", text="Thành tiền")

        self.tree_cart.column(
            "pcode", width=80, minwidth=65, anchor="center"
        )
        self.tree_cart.column(
            "pname", width=300, minwidth=150
        )
        self.tree_cart.column(
            "qty", width=100, minwidth=75, anchor="center"
        )
        self.tree_cart.column(
            "price", width=150, minwidth=100, anchor="e"
        )
        self.tree_cart.column(
            "total", width=170, minwidth=110, anchor="e"
        )

        cart_scrollbar = ttk.Scrollbar(
            table_frame,
            orient="vertical",
            command=self.tree_cart.yview
        )

        cart_horizontal_scrollbar = ttk.Scrollbar(
            table_frame,
            orient="horizontal",
            command=self.tree_cart.xview
        )

        self.tree_cart.configure(
            yscrollcommand=cart_scrollbar.set,
            xscrollcommand=cart_horizontal_scrollbar.set
        )

        table_frame.grid_rowconfigure(0, weight=1)
        table_frame.grid_columnconfigure(0, weight=1)

        self.tree_cart.grid(
            row=0,
            column=0,
            sticky="nsew"
        )

        cart_scrollbar.grid(
            row=0,
            column=1,
            sticky="ns"
        )

        cart_horizontal_scrollbar.grid(
            row=1,
            column=0,
            sticky="ew"
        )

        # ==========================================
        # NÚT GIỎ HÀNG VÀ TỔNG TIỀN
        # ==========================================

        cart_button_frame = ctk.CTkFrame(
            cart_frame,
            fg_color=card_color
        )
        cart_button_frame.pack(
            fill="x",
            padx=12,
            pady=(5, 12)
        )

        ctk.CTkButton(
            cart_button_frame,
            text="Xóa sản phẩm",
            width=120,
            height=32,
            corner_radius=8,
            fg_color="#E74C3C",
            hover_color="#C0392B",
            command=self.remove_from_cart
        ).pack(side="left", padx=5)

        ctk.CTkButton(
            cart_button_frame,
            text="Xóa giỏ hàng",
            width=120,
            height=32,
            corner_radius=8,
            fg_color="#85859B",
            hover_color="#6B6B80",
            command=self.clear_cart
        ).pack(side="left", padx=5)

        self.lbl_total = ctk.CTkLabel(
            cart_button_frame,
            text="Tổng tiền: 0 VNĐ",
            font=("Arial", 20, "bold"),
            text_color=purple_color
        )
        self.lbl_total.pack(
            side="right",
            padx=15
        )

        # ==========================================
        # THANH TOÁN
        # ==========================================

        payment_frame = ctk.CTkFrame(
            self.parent,
            fg_color=card_color,
            corner_radius=12,
            border_width=1,
            border_color=border_color
        )
        payment_frame.pack(
            fill="x",
            padx=25,
            pady=6
        )

        ctk.CTkLabel(
            payment_frame,
            text="Thanh toán",
            font=("Arial", 16, "bold"),
            text_color=text_color
        ).grid(
            row=0,
            column=0,
            columnspan=6,
            padx=15,
            pady=(10, 2),
            sticky="w"
        )

        ctk.CTkLabel(
            payment_frame,
            text="Tiền khách đưa:",
            text_color=text_color
        ).grid(
            row=1,
            column=0,
            padx=12,
            pady=10
        )

        self.txt_payment = ctk.CTkEntry(
            payment_frame,
            width=160,
            height=34,
            corner_radius=8,
            border_color=border_color,
            fg_color="#FAFAFC",
            text_color="#000000",
            placeholder_text="Nhập số tiền"
        )
        self.txt_payment.grid(
            row=1,
            column=1,
            padx=8,
            pady=10
        )

        ctk.CTkLabel(
            payment_frame,
            text="Tiền thừa:",
            text_color=text_color
        ).grid(
            row=1,
            column=2,
            padx=12,
            pady=10
        )

        self.lbl_change = ctk.CTkLabel(
            payment_frame,
            text="0 VNĐ",
            font=("Arial", 16, "bold"),
            text_color=purple_color
        )
        self.lbl_change.grid(
            row=1,
            column=3,
            padx=8,
            pady=10
        )

        ctk.CTkButton(
            payment_frame,
            text="Tính tiền thừa",
            width=130,
            height=34,
            corner_radius=8,
            fg_color="#4385F5",
            hover_color="#2867D6",
            command=self.calculate_change
        ).grid(
            row=1,
            column=4,
            padx=10,
            pady=10
        )

        ctk.CTkButton(
            payment_frame,
            text="THANH TOÁN",
            width=155,
            height=38,
            corner_radius=8,
            fg_color=purple_color,
            hover_color=hover_color,
            font=("Arial", 13, "bold"),
            command=self.checkout
        ).grid(
            row=1,
            column=5,
            padx=(10, 15),
            pady=10
        )

        # ==========================================
        # LỊCH SỬ HÓA ĐƠN
        # ==========================================

        history_frame = ctk.CTkFrame(
            self.parent,
            fg_color=card_color,
            corner_radius=12,
            border_width=1,
            border_color=border_color
        )
        history_frame.pack(
            fill="x",
            padx=25,
            pady=(6, 16)
        )

        history_header = ctk.CTkFrame(
            history_frame,
            fg_color=card_color
        )
        history_header.pack(
            fill="x",
            padx=15,
            pady=(10, 3)
        )

        ctk.CTkLabel(
            history_header,
            text="Lịch sử hóa đơn",
            font=("Arial", 17, "bold"),
            text_color=text_color
        ).pack(side="left")

        # ==========================================
        # BẢNG LỊCH SỬ
        # ==========================================

        history_table_frame = ctk.CTkFrame(
            history_frame,
            fg_color=card_color
        )
        history_table_frame.pack(
            fill="x",
            padx=12,
            pady=5
        )

        history_columns = (
            "transno",
            "customer",
            "cashier",
            "total"
        )

        self.tree_history = ttk.Treeview(
            history_table_frame,
            columns=history_columns,
            show="headings",
            height=4,
            style="Cash.Treeview"
        )

        self.tree_history.heading("transno", text="Mã hóa đơn")
        self.tree_history.heading("customer", text="Khách hàng")
        self.tree_history.heading("cashier", text="Nhân viên")
        self.tree_history.heading("total", text="Tổng tiền")

        self.tree_history.column(
            "transno", width=140, minwidth=100, anchor="center"
        )
        self.tree_history.column(
            "customer", width=250, minwidth=140
        )
        self.tree_history.column(
            "cashier", width=200, minwidth=120
        )
        self.tree_history.column(
            "total", width=180, minwidth=120, anchor="e"
        )

        history_scrollbar = ttk.Scrollbar(
            history_table_frame,
            orient="vertical",
            command=self.tree_history.yview
        )

        self.tree_history.configure(
            yscrollcommand=history_scrollbar.set
        )

        self.tree_history.pack(
            side="left",
            fill="both",
            expand=True
        )

        history_scrollbar.pack(
            side="right",
            fill="y"
        )

        # ==========================================
        # NÚT CHI TIẾT VÀ XUẤT HÓA ĐƠN
        # ==========================================

        history_button_frame = ctk.CTkFrame(
            history_frame,
            fg_color=card_color
        )
        history_button_frame.pack(
            fill="x",
            padx=12,
            pady=(3, 12)
        )

        ctk.CTkButton(
            history_button_frame,
            text="Xem chi tiết hóa đơn",
            width=180,
            height=34,
            corner_radius=8,
            fg_color="#4385F5",
            hover_color="#2867D6",
            command=self.view_invoice_detail
        ).pack(
            side="left",
            padx=5
        )

        ctk.CTkButton(
            history_button_frame,
            text="Xuất hóa đơn",
            width=160,
            height=34,
            corner_radius=8,
            fg_color=purple_color,
            hover_color=hover_color,
            font=("Arial", 12, "bold"),
            command=self.export_selected_invoice
        ).pack(
            side="left",
            padx=5
        )

    # LOAD KHÁCH HÀNG

    def load_customers(self):
        connection = None
        cursor = None

        try:
            connection = connect_database()
            cursor = connection.cursor()

            cursor.execute("""
                SELECT id, name, phone
                FROM tbCustomer
                ORDER BY id
            """)

            rows = cursor.fetchall()

            self.customer_data = {}
            values = []

            for row in rows:
                customer_id = row[0]
                name = row[1]
                phone = row[2]

                display = str(customer_id) + " - " + str(name)

                if phone:
                    display = display + " - " + str(phone)

                values.append(display)
                self.customer_data[display] = customer_id

            if len(values) == 0:
                self.cbo_customer.configure(values=[""])
                self.cbo_customer.set("")
                return

            self.cbo_customer.configure(values=values)
            self.cbo_customer.set(values[0])

            # Tự động tải dịch vụ của khách hàng đầu tiên.
            self.customer_selected(values[0])

        except Exception as e:
            messagebox.showerror(
                "Lỗi",
                "Không thể tải danh sách khách hàng.\n\n" + str(e)
            )

        finally:
            if cursor is not None:
                cursor.close()

            if connection is not None:
                connection.close()


    def refresh_cart(self):
        for item in self.tree_cart.get_children():
            self.tree_cart.delete(item)

        total_amount = 0

        for item in self.cart:
            self.tree_cart.insert(
                "",
                "end",
                values=(
                    item["pcode"],
                    item["pname"],
                    item["qty"],
                    "{:,.0f}".format(item["price"]),
                    "{:,.0f}".format(item["total"])
                )
            )

            total_amount += item["total"]

        self.lbl_total.configure(
            text="{:,.0f} VNĐ".format(total_amount)
        )

    # KHI CHỌN KHÁCH HÀNG

    def customer_selected(self, selected_value=None):
        if selected_value is None:
            selected_value = self.cbo_customer.get()

        if selected_value not in self.customer_data:
            return

        customer_id = self.customer_data[selected_value]

        self.load_customer_services(customer_id)

    def load_customer_services(self, customer_id):
        connection = None
        cursor = None

        try:
            connection = connect_database()
            cursor = connection.cursor()

            sql = """
                SELECT
                    ps.PetServiceID,
                    ps.PetID,
                    p.PetName,
                    s.ServiceName,
                    ps.Price
                FROM tbPetService ps
                INNER JOIN tbPet p
                    ON ps.PetID = p.PetID
                INNER JOIN tbService s
                    ON ps.ServiceID = s.ServiceID
                WHERE p.CustomerID = ?
                  AND ps.Price IS NOT NULL
                  AND NOT EXISTS (
                      SELECT 1
                      FROM tbCash c
                      WHERE c.pcode =
                          N'DV' + CAST(
                              ps.PetServiceID AS NVARCHAR(20)
                          )
                  )
                ORDER BY ps.ServiceDate DESC
            """

            cursor.execute(sql, (customer_id,))
            rows = cursor.fetchall()

            # Giữ các sản phẩm bán lẻ trong giỏ.
            new_cart = []

            for item in self.cart:
                if "service_id" not in item:
                    new_cart.append(item)

            self.cart = new_cart

            # Thêm dịch vụ chưa thanh toán của khách vào giỏ.
            for row in rows:
                pet_service_id = row[0]
                pet_id = row[1]
                pet_name = row[2]
                service_name = row[3]
                price = float(row[4])

                item = {
                    "service_id": pet_service_id,
                    "pet_id": pet_id,
                    "pcode": "DV" + str(pet_service_id),
                    "pname": str(pet_name) + " - " + str(service_name),
                    "qty": 1,
                    "price": price,
                    "total": price
                }

                self.cart.append(item)

            self.refresh_cart()

        except Exception as e:
            messagebox.showerror(
                "Lỗi",
                "Không thể tải dịch vụ của khách hàng.\n\n" + str(e)
            )

        finally:
            if cursor is not None:
                cursor.close()

            if connection is not None:
                connection.close()

    def load_ready_services(self, customer_id):
        # Xóa dịch vụ tự động của khách hàng trước.
        # Giữ nguyên các sản phẩm bán lẻ trong giỏ.
        new_cart = []

        for item in self.cart:
            if "service_id" not in item:
                new_cart.append(item)

        self.cart = new_cart

        if customer_id is None:
            self.refresh_cart()
            return

        conn = None
        cursor = None

        try:
            conn = connect_database()
            cursor = conn.cursor()

            sql = """
                SELECT
                    b.BoardingID,
                    ps.PetServiceID,
                    ps.PetID,
                    p.PetName,
                    s.ServiceName,
                    ps.Price
                FROM tbPetBoarding b
                INNER JOIN tbPet p
                    ON b.PetID = p.PetID
                INNER JOIN tbPetService ps
                    ON ps.PetID = p.PetID
                INNER JOIN tbService s
                    ON ps.ServiceID = s.ServiceID
                WHERE b.CustomerID = ?
                  AND LTRIM(RTRIM(b.Status)) = N'Sẵn sàng'
                  AND LTRIM(RTRIM(ps.Status)) = N'Đã hoàn thành'
                  AND ps.Price IS NOT NULL
                  AND NOT EXISTS (
                      SELECT 1
                      FROM tbCash c
                      WHERE c.pcode =
                          N'DV' + CAST(
                              ps.PetServiceID AS NVARCHAR(20)
                          )
                  )
            """

            cursor.execute(sql, (customer_id,))
            rows = cursor.fetchall()

            if len(rows) == 0:
                print(
                    "Không tìm thấy dịch vụ đủ điều kiện cho khách:",
                    customer_id
                )

            for row in rows:
                boarding_id = row[0]
                pet_service_id = row[1]
                pet_id = row[2]
                pet_name = str(row[3])
                service_name = str(row[4])
                price = float(row[5])

                item = {
                    "boarding_id": boarding_id,
                    "service_id": pet_service_id,
                    "pet_id": pet_id,
                    "pcode": "DV" + str(pet_service_id),
                    "pname": pet_name + " - " + service_name,
                    "qty": 1,
                    "price": price,
                    "total": price
                }

                self.cart.append(item)

        except Exception as e:
            messagebox.showerror(
                "Lỗi tải dịch vụ",
                str(e)
            )

        finally:
            if cursor is not None:
                cursor.close()

            if conn is not None:
                conn.close()

        self.refresh_cart()

    def load_customer_boarding(self, customer_id):

        for item in self.tree_boarding.get_children():
            self.tree_boarding.delete(item)

        if customer_id is None:
            self.lbl_boarding_status.configure(
                text="Vui lòng chọn khách hàng",
                text_color="#85859B"
            )
            return

        connection = None

        try:
            connection = connect_database()
            cursor = connection.cursor()

            sql = """
                   SELECT
                       b.BoardingID,
                       p.PetName,
                       b.Status,
                       b.CheckInDate,
                       b.ExpectedDate,
                       b.Area,
                       b.Note
                   FROM tbPetBoarding b
                   INNER JOIN tbPet p
                       ON b.PetID = p.PetID
                   WHERE b.CustomerID = ?
                   ORDER BY b.BoardingID DESC
               """

            cursor.execute(sql, (customer_id,))
            rows = cursor.fetchall()

            count = 0
            count_booking = 0
            count_staying = 0
            count_returned = 0

            for row in rows:

                boarding_id = row[0]
                pet_name = row[1]
                status = str(row[2] or "").strip()
                checkin_date = row[3]
                expected_date = row[4]
                area = row[5] or ""
                note = row[6] or ""

                # Định dạng ngày
                if checkin_date is not None:
                    checkin_date = checkin_date.strftime("%d/%m/%Y")
                else:
                    checkin_date = ""

                if expected_date is not None:
                    expected_date = expected_date.strftime("%d/%m/%Y")
                else:
                    expected_date = ""

                # Phân loại trạng thái
                if status in ("Đã đặt", "Đã đặt lịch"):
                    count_booking += 1

                elif status in (
                        "Đang gửi",
                        "Đang lưu trú",
                        "Đang tắm"
                ):
                    count_staying += 1

                elif status in ("Đã trả", "Đã thanh toán"):
                    count_returned += 1

                self.tree_boarding.insert(
                    "",
                    "end",
                    values=(
                        boarding_id,
                        pet_name,
                        status,
                        checkin_date,
                        expected_date,
                        area,
                        note
                    )
                )

                count += 1

            self.lbl_boarding_status.configure(
                text=(
                        "Tổng lượt lưu trú: " + str(count)
                        + "  |  Đã đặt: " + str(count_booking)
                        + "  |  Đang gửi: " + str(count_staying)
                        + "  |  Đã trả: " + str(count_returned)
                ),
                text_color="#202044"
            )

            if count == 0:
                self.lbl_boarding_status.configure(
                    text="Khách hàng này chưa có thông tin lưu trú.",
                    text_color="#85859B"
                )

            cursor.close()
            connection.close()

        except Exception as e:

            if connection is not None:
                try:
                    connection.close()
                except:
                    pass

            messagebox.showerror(
                "Lỗi",
                "Không thể tải thông tin lưu trú.\n\n" + str(e)
            )

    # LOAD SẢN PHẨM

    def load_products(self):

        try:

            connection = connect_database()
            cursor = connection.cursor()

            sql = """
                SELECT
                    pcode,
                    pname
                FROM tbProduct
                ORDER BY pcode
            """

            cursor.execute(sql)

            rows = cursor.fetchall()

            self.product_data = {}

            values = []

            for row in rows:

                pcode = row[0]
                pname = row[1]

                display = (
                    str(pcode)
                    + " - "
                    + str(pname)
                )

                values.append(display)

                self.product_data[display] = pcode

            if len(values) == 0:

                values.append("")

            self.cbo_product.configure(
                values=values
            )

            self.cbo_product.set(
                values[0]
            )

            if values[0] != "":

                self.product_selected(
                    values[0]
                )

            cursor.close()
            connection.close()

        except Exception as e:

            messagebox.showerror(
                "Lỗi",
                "Không thể tải danh sách sản phẩm.\n\n"
                + str(e)
            )

    # CHỌN SẢN PHẨM

    def product_selected(self, selected_value):

        if selected_value == "":

            return

        if selected_value not in self.product_data:

            return

        pcode = self.product_data[
            selected_value
        ]

        try:

            connection = connect_database()
            cursor = connection.cursor()

            sql = """
                SELECT
                    pname,
                    pqty,
                    pprice
                FROM tbProduct
                WHERE pcode = ?
            """

            cursor.execute(
                sql,
                (pcode,)
            )

            row = cursor.fetchone()

            cursor.close()
            connection.close()

            if row:

                price = float(row[2])
                stock = row[1]

                self.txt_price.configure(
                    state="normal"
                )

                self.txt_price.delete(
                    0,
                    "end"
                )

                self.txt_price.insert(
                    0,
                    "{:,.0f}".format(price)
                )

                self.txt_price.configure(
                    state="disabled"
                )

                self.lbl_stock.configure(
                    text=str(stock)
                )

        except Exception as e:

            messagebox.showerror(
                "Lỗi",
                "Không thể lấy thông tin sản phẩm.\n\n"
                + str(e)
            )

    # THÊM VÀO GIỎ

    def add_to_cart(self):

        selected_product = self.cbo_product.get()

        if selected_product == "":

            messagebox.showwarning(
                "Thông báo",
                "Vui lòng chọn sản phẩm."
            )

            return

        if selected_product not in self.product_data:

            messagebox.showwarning(
                "Thông báo",
                "Sản phẩm không hợp lệ."
            )

            return

        try:

            quantity = int(
                self.txt_quantity.get()
            )

        except:

            messagebox.showwarning(
                "Thông báo",
                "Số lượng phải là số nguyên."
            )

            return

        if quantity <= 0:

            messagebox.showwarning(
                "Thông báo",
                "Số lượng phải lớn hơn 0."
            )

            return

        pcode = self.product_data[
            selected_product
        ]

        try:

            connection = connect_database()
            cursor = connection.cursor()

            sql = """
                SELECT
                    pname,
                    pqty,
                    pprice
                FROM tbProduct
                WHERE pcode = ?
            """

            cursor.execute(
                sql,
                (pcode,)
            )

            row = cursor.fetchone()

            cursor.close()
            connection.close()

            if row is None:

                messagebox.showerror(
                    "Lỗi",
                    "Không tìm thấy sản phẩm."
                )

                return

            pname = row[0]
            stock = int(row[1])
            price = float(row[2])

            existing_quantity = 0

            for item in self.cart:

                if item["pcode"] == pcode:

                    existing_quantity = item["qty"]

            new_quantity = (
                existing_quantity
                + quantity
            )

            if new_quantity > stock:

                messagebox.showwarning(
                    "Không đủ hàng",
                    "Sản phẩm chỉ còn "
                    + str(stock)
                    + " sản phẩm."
                )

                return

            found = False

            for item in self.cart:

                if item["pcode"] == pcode:

                    item["qty"] = new_quantity

                    item["total"] = (
                        new_quantity
                        * price
                    )

                    found = True

            if found == False:

                new_item = {
                    "pcode": pcode,
                    "pname": pname,
                    "qty": quantity,
                    "price": price,
                    "total": quantity * price
                }

                self.cart.append(
                    new_item
                )

            self.refresh_cart()

        except Exception as e:

            messagebox.showerror(
                "Lỗi",
                "Không thể thêm sản phẩm.\n\n"
                + str(e)
            )

    # HIỂN THỊ GIỎ HÀNG

    def refresh_cart(self):
        for item in self.tree_cart.get_children():
            self.tree_cart.delete(item)

        total_amount = 0

        for item in self.cart:
            self.tree_cart.insert(
                "",
                "end",
                values=(
                    item["pcode"],
                    item["pname"],
                    item["qty"],
                    "{:,.0f}".format(float(item["price"])),
                    "{:,.0f}".format(float(item["total"]))
                )
            )

            total_amount += float(item["total"])

        self.lbl_total.configure(
            text="Tổng tiền: {:,.0f} VNĐ".format(total_amount)
        )

    # XÓA SẢN PHẨM KHỎI GIỎ

    def remove_from_cart(self):
        selected = self.tree_cart.selection()

        if len(selected) == 0:
            messagebox.showwarning(
                "Thông báo",
                "Vui lòng chọn sản phẩm cần xóa."
            )
            return

        item = self.tree_cart.item(selected[0])
        pcode = str(item["values"][0])

        new_cart = []

        for cart_item in self.cart:
            if str(cart_item["pcode"]) != pcode:
                new_cart.append(cart_item)

        self.cart = new_cart
        self.refresh_cart()

    # XÓA GIỎ HÀNG

    def clear_cart(self):

        if len(self.cart) == 0:

            return

        answer = messagebox.askyesno(
            "Xác nhận",
            "Bạn có chắc muốn xóa toàn bộ giỏ hàng?"
        )

        if answer:

            self.cart = []

            self.refresh_cart()

            self.txt_payment.delete(
                0,
                "end"
            )

            self.lbl_change.configure(
                text="0 VNĐ"
            )

    # TÍNH TIỀN THỪA

    def calculate_change(self):

        if len(self.cart) == 0:

            messagebox.showwarning(
                "Thông báo",
                "Giỏ hàng đang trống."
            )

            return

        total_money = 0

        for item in self.cart:

            total_money = (
                total_money
                + item["total"]
            )

        try:

            payment = float(
                self.txt_payment.get()
            )

        except:

            messagebox.showwarning(
                "Thông báo",
                "Tiền khách đưa phải là số."
            )

            return

        if payment < total_money:

            messagebox.showwarning(
                "Thông báo",
                "Số tiền khách đưa chưa đủ."
            )

            return

        change = payment - total_money

        self.lbl_change.configure(
            text="{:,.0f} VNĐ".format(
                change
            )
        )

    # TẠO MÃ HÓA ĐƠN

    def generate_transaction_number(self):

        try:

            connection = connect_database()
            cursor = connection.cursor()

            sql = """
                SELECT transno
                FROM tbCash
                WHERE transno IS NOT NULL
            """

            cursor.execute(sql)

            rows = cursor.fetchall()

            cursor.close()
            connection.close()

            max_number = 0

            for row in rows:

                transno = row[0]

                if transno is None:

                    continue

                transno = str(
                    transno
                )

                if transno.startswith("HD"):

                    number_text = transno[2:]

                    try:

                        number = int(
                            number_text
                        )

                        if number > max_number:

                            max_number = number

                    except:

                        pass

            new_number = (
                max_number + 1
            )

            return "HD{:04d}".format(
                new_number
            )

        except Exception as e:

            messagebox.showerror(
                "Lỗi",
                "Không thể tạo mã hóa đơn.\n\n"
                + str(e)
            )

            return None

    # THANH TOÁN

    def checkout(self):

        if len(self.cart) == 0:

            messagebox.showwarning(
                "Thông báo",
                "Giỏ hàng đang trống."
            )

            return

        customer_display = (
            self.cbo_customer.get()
        )

        if customer_display == "":

            messagebox.showwarning(
                "Thông báo",
                "Vui lòng chọn khách hàng."
            )

            return

        if customer_display not in self.customer_data:

            messagebox.showwarning(
                "Thông báo",
                "Khách hàng không hợp lệ."
            )

            return

        customer_id = (
            self.customer_data[
                customer_display
            ]
        )

        total_money = 0

        for item in self.cart:

            total_money = (
                total_money
                + item["total"]
            )

        try:

            payment = float(
                self.txt_payment.get()
            )

        except:

            messagebox.showwarning(
                "Thông báo",
                "Vui lòng nhập số tiền khách đưa."
            )

            return

        if payment < total_money:

            messagebox.showwarning(
                "Thông báo",
                "Tiền khách đưa chưa đủ để thanh toán."
            )

            return

        transaction_number = (
            self.generate_transaction_number()
        )

        if transaction_number is None:

            return

        cashier = Session.full_name

        if cashier == "":

            cashier = Session.username

        connection = None

        try:

            connection = connect_database()

            cursor = connection.cursor()

            # =================================================
            # THÊM CHI TIẾT HÓA ĐƠN
            # =================================================

            for item in self.cart:

                sql_insert = """
                    INSERT INTO tbCash
                    (
                        transno,
                        pcode,
                        pname,
                        qty,
                        price,
                        total,
                        cid,
                        cashier
                    )
                    VALUES
                    (
                        ?, ?, ?, ?, ?, ?, ?, ?
                    )
                """

                cursor.execute(
                    sql_insert,
                    (
                        transaction_number,
                        str(item["pcode"]),
                        item["pname"],
                        item["qty"],
                        item["price"],
                        item["total"],
                        str(customer_id),
                        cashier
                    )
                )

                # =================================================
                # TRỪ TỒN KHO
                # =================================================

                # Chỉ trừ tồn kho đối với sản phẩm bán lẻ
                if "service_id" not in item:
                    cursor.execute("""
                        UPDATE tbProduct
                        SET pqty = pqty - ?
                        WHERE pcode = ? AND pqty >= ?
                    """, (
                        item["qty"],
                        item["pcode"],
                        item["qty"]
                    ))

                    if cursor.rowcount == 0:
                        raise Exception(
                            "Sản phẩm không tồn tại hoặc không đủ tồn kho: "
                            + str(item["pname"])
                        )

                if cursor.rowcount == 0:

                    connection.rollback()

                    messagebox.showerror(
                        "Lỗi",
                        "Sản phẩm "
                        + item["pname"]
                        + " không đủ số lượng."
                    )

                    cursor.close()
                    connection.close()

                    return

            # Cập nhật trạng thái lưu trú sau khi thanh toán dịch vụ
            for item in self.cart:

                boarding_id = item.get("boarding_id")
                service_id = item.get("service_id")

                if boarding_id is not None and service_id is not None:
                    sql_update = """
                        UPDATE tbPetBoarding
                        SET Status = N'Đã trả'
                        WHERE BoardingID = ?
                          AND Status = N'Sẵn sàng'
                    """

                    cursor.execute(
                        sql_update,
                        (boarding_id,)
                    )

            # Lưu hóa đơn và trạng thái trong cùng giao dịch

            connection.commit()

            cursor.close()
            connection.close()

            change = (
                payment
                - total_money
            )

            messagebox.showinfo(
                "Thanh toán thành công",
                "Mã hóa đơn: "
                + transaction_number
                + "\n\n"
                + "Tổng tiền: "
                + "{:,.0f}".format(
                    total_money
                )
                + " VNĐ"
                + "\n"
                + "Khách đưa: "
                + "{:,.0f}".format(
                    payment
                )
                + " VNĐ"
                + "\n"
                + "Tiền thừa: "
                + "{:,.0f}".format(
                    change
                )
                + " VNĐ"
            )

            self.cart = []

            self.refresh_cart()

            self.txt_payment.delete(
                0,
                "end"
            )

            self.lbl_change.configure(
                text="0 VNĐ"
            )

            self.load_products()

            self.load_invoice_history()

        except Exception as e:

            if connection is not None:

                try:

                    connection.rollback()

                except:

                    pass

                try:

                    connection.close()

                except:

                    pass

            messagebox.showerror(
                "Lỗi thanh toán",
                "Không thể lưu hóa đơn.\n\n"
                + str(e)
            )

    # LOAD LỊCH SỬ HÓA ĐƠN

    def load_invoice_history(self):

        for item in self.tree_history.get_children():

            self.tree_history.delete(item)

        try:

            connection = connect_database()
            cursor = connection.cursor()

            sql = """
                SELECT
                    c.transno,
                    c.cid,
                    ISNULL(
                        cust.name,
                        N'Khách lẻ'
                    ),
                    c.cashier,
                    SUM(c.total)
                FROM tbCash c
                LEFT JOIN tbCustomer cust
                    ON c.cid =
                       CAST(
                           cust.id AS NVARCHAR(20)
                       )
                GROUP BY
                    c.transno,
                    c.cid,
                    cust.name,
                    c.cashier
                ORDER BY c.transno DESC
            """

            cursor.execute(sql)

            rows = cursor.fetchall()

            for row in rows:

                self.tree_history.insert(
                    "",
                    "end",
                    values=(
                        row[0],
                        row[2],
                        row[3],
                        "{:,.0f}".format(
                            float(row[4])
                        )
                    )
                )

            cursor.close()
            connection.close()

        except Exception as e:

            messagebox.showerror(
                "Lỗi",
                "Không thể tải lịch sử hóa đơn.\n\n"
                + str(e)
            )

    # XEM CHI TIẾT HÓA ĐƠN

    def view_invoice_detail(self):

        selected = (
            self.tree_history.selection()
        )

        if len(selected) == 0:

            messagebox.showwarning(
                "Thông báo",
                "Vui lòng chọn hóa đơn."
            )

            return

        item = self.tree_history.item(
            selected[0]
        )

        transno = item["values"][0]

        try:

            connection = connect_database()
            cursor = connection.cursor()

            sql = """
                SELECT
                    pname,
                    qty,
                    price,
                    total
                FROM tbCash
                WHERE transno = ?
                ORDER BY cashid
            """

            cursor.execute(
                sql,
                (transno,)
            )

            rows = cursor.fetchall()

            cursor.close()
            connection.close()

            detail = ""

            detail = (
                detail
                + "HÓA ĐƠN: "
                + str(transno)
            )

            detail = (
                detail
                + "\n"
                + "-" * 50
                + "\n"
            )

            total_money = 0

            for row in rows:

                pname = row[0]
                quantity = row[1]
                price = float(row[2])
                total = float(row[3])

                detail = (
                    detail
                    + str(pname)
                    + "\n"
                )

                detail = (
                    detail
                    + "  SL: "
                    + str(quantity)
                )

                detail = (
                    detail
                    + " | Đơn giá: "
                    + "{:,.0f}".format(price)
                )

                detail = (
                    detail
                    + " | Thành tiền: "
                    + "{:,.0f}".format(total)
                )

                detail = (
                    detail
                    + "\n"
                )

                total_money = (
                    total_money
                    + total
                )

            detail = (
                detail
                + "-"
                * 50
                + "\n"
            )

            detail = (
                detail
                + "TỔNG TIỀN: "
                + "{:,.0f}".format(
                    total_money
                )
                + " VNĐ"
            )

            messagebox.showinfo(
                "Chi tiết hóa đơn",
                detail
            )

        except Exception as e:

            messagebox.showerror(
                "Lỗi",
                "Không thể xem chi tiết hóa đơn.\n\n"
                + str(e)
            )


    # XUẤT HÓA ĐƠN

    def export_selected_invoice(self):

        selected = self.tree_history.selection()

        if len(selected) == 0:

            messagebox.showwarning(
                "Thông báo",
                "Vui lòng chọn hóa đơn cần xuất."
            )

            return

        item = self.tree_history.item(
            selected[0]
        )

        transno = str(item["values"][0])

        export_window = ctk.CTkToplevel(self.parent)
        export_window.title("Xuất hóa đơn")
        export_window.geometry("430x330")
        export_window.resizable(False, False)
        export_window.transient(self.parent.winfo_toplevel())
        export_window.grab_set()

        title = ctk.CTkLabel(
            export_window,
            text="XUẤT HÓA ĐƠN",
            font=("Arial", 22, "bold")
        )

        title.pack(pady=(25, 5))

        invoice_label = ctk.CTkLabel(
            export_window,
            text="Mã hóa đơn: " + transno,
            font=("Arial", 14, "bold")
        )

        invoice_label.pack(pady=(0, 20))

        self.export_window = export_window
        self.export_transno = transno

        txt_button = ctk.CTkButton(
            export_window,
            text="XUẤT FILE TXT",
            width=260,
            height=40,
            command=self.export_txt_button
        )

        txt_button.pack(pady=7)

        pdf_button = ctk.CTkButton(
            export_window,
            text="XUẤT FILE PDF",
            width=260,
            height=40,
            command=self.export_pdf_button
        )

        pdf_button.pack(pady=7)

        docx_button = ctk.CTkButton(
            export_window,
            text="XUẤT FILE DOCX",
            width=260,
            height=40,
            command=self.export_docx_button
        )

        docx_button.pack(pady=7)

        note = ctk.CTkLabel(
            export_window,
            text="File sẽ được lưu trong thư mục invoices.",
            font=("Arial", 12)
        )

        note.pack(pady=(10, 0))

    def export_txt_button(self):

        self.export_invoice_format(
            self.export_transno,
            "txt",
            self.export_window
        )

    def export_pdf_button(self):

        self.export_invoice_format(
            self.export_transno,
            "pdf",
            self.export_window
        )

    def export_docx_button(self):

        self.export_invoice_format(
            self.export_transno,
            "docx",
            self.export_window
        )

    # LẤY DỮ LIỆU HÓA ĐƠN

    def get_invoice_data(self, transno):

        connection = None

        try:

            connection = connect_database()
            cursor = connection.cursor()

            sql = """
                SELECT
                    c.transno,
                    c.pcode,
                    c.pname,
                    c.qty,
                    c.price,
                    c.total,
                    c.cid,
                    ISNULL(
                        cust.name,
                        N'Khách lẻ'
                    ) AS CustomerName,
                    c.cashier
                FROM tbCash c
                LEFT JOIN tbCustomer cust
                    ON c.cid = CAST(cust.id AS NVARCHAR(20))
                WHERE c.transno = ?
                ORDER BY c.cashid
            """

            cursor.execute(
                sql,
                (transno,)
            )

            rows = cursor.fetchall()

            cursor.close()
            connection.close()
            connection = None

            if len(rows) == 0:

                return None

            invoice = {
                "transno": str(transno),
                "customer": str(rows[0][7]),
                "cashier": str(rows[0][8]),
                "items": [],
                "total": 0
            }

            for row in rows:

                item = {
                    "pcode": str(row[1]),
                    "pname": str(row[2]),
                    "qty": int(row[3]),
                    "price": float(row[4]),
                    "total": float(row[5])
                }

                invoice["items"].append(item)

                invoice["total"] = (
                    invoice["total"]
                    + item["total"]
                )

            return invoice

        except Exception as e:

            if connection is not None:

                try:
                    connection.close()
                except:
                    pass

            messagebox.showerror(
                "Lỗi",
                "Không thể lấy dữ liệu hóa đơn.\n\n"
                + str(e)
            )

            return None

    # TẠO THƯ MỤC LƯU HÓA ĐƠN

    def get_invoice_folder(self):

        folder = os.path.join(
            os.getcwd(),
            "invoices"
        )

        if not os.path.exists(folder):

            os.makedirs(folder)

        return folder

    # TẠO QR CODE

    def create_invoice_qr(self, invoice):

        folder = self.get_invoice_folder()

        qr_file = os.path.join(
            folder,
            invoice["transno"] + "_QR.png"
        )

        qr_text = (
            "PET SHOP MANAGEMENT SYSTEM\n"
            "HOA DON: "
            + invoice["transno"]
            + "\n"
            + "KHACH HANG: "
            + invoice["customer"]
            + "\n"
            + "NHAN VIEN: "
            + invoice["cashier"]
            + "\n"
            + "TONG TIEN: "
            + "{:,.0f}".format(invoice["total"])
            + " VND"
        )

        qr = qrcode.QRCode(
            version=1,
            error_correction=qrcode.constants.ERROR_CORRECT_M,
            box_size=10,
            border=4
        )

        qr.add_data(qr_text)
        qr.make(fit=True)

        image = qr.make_image()
        image.save(qr_file)

        return qr_file

    # XUẤT TXT

    def export_invoice_txt(self, invoice):

        folder = self.get_invoice_folder()

        file_path = os.path.join(
            folder,
            invoice["transno"] + ".txt"
        )

        self.create_invoice_qr(invoice)

        lines = []

        lines.append("=" * 60)
        lines.append("PET SHOP MANAGEMENT SYSTEM")
        lines.append("HÓA ĐƠN BÁN HÀNG")
        lines.append("=" * 60)
        lines.append("Mã hóa đơn: " + invoice["transno"])
        lines.append("Khách hàng: " + invoice["customer"])
        lines.append("Nhân viên: " + invoice["cashier"])
        lines.append("-" * 60)
        lines.append("Sản phẩm")
        lines.append("-" * 60)

        for item in invoice["items"]:

            lines.append(
                str(item["pcode"])
                + " - "
                + item["pname"]
            )

            lines.append(
                "  SL: "
                + str(item["qty"])
                + " | Đơn giá: "
                + "{:,.0f}".format(item["price"])
                + " VNĐ"
                + " | Thành tiền: "
                + "{:,.0f}".format(item["total"])
                + " VNĐ"
            )

        lines.append("-" * 60)
        lines.append(
            "TỔNG TIỀN: {:,.0f} VNĐ".format(
                invoice["total"]
            )
        )
        lines.append("-" * 60)
        lines.append("QR Code: " + invoice["transno"] + "_QR.png")
        lines.append("=" * 60)
        lines.append("Cảm ơn quý khách!")
        lines.append("=" * 60)

        with open(
            file_path,
            "w",
            encoding="utf-8"
        ) as file:

            file.write("\n".join(lines))

        return file_path

    # TÌM FONT TIẾNG VIỆT CHO PDF

    def register_pdf_font(self):

        font_paths = [
            "C:/Windows/Fonts/arial.ttf",
            "C:/Windows/Fonts/ARIAL.TTF",
            "C:/Windows/Fonts/calibri.ttf",
            "C:/Windows/Fonts/Calibri.ttf",
            "C:/Windows/Fonts/tahoma.ttf",
            "C:/Windows/Fonts/Tahoma.ttf"
        ]

        for font_path in font_paths:

            if os.path.exists(font_path):

                try:

                    pdfmetrics.registerFont(
                        TTFont(
                            "PetShopUnicode",
                            font_path
                        )
                    )

                    return "PetShopUnicode"

                except:

                    pass

        return "Helvetica"

    # XUẤT PDF

    def export_invoice_pdf(self, invoice):

        folder = self.get_invoice_folder()

        file_path = os.path.join(
            folder,
            invoice["transno"] + ".pdf"
        )

        qr_file = self.create_invoice_qr(invoice)

        pdf_font = self.register_pdf_font()

        document = SimpleDocTemplate(
            file_path,
            pagesize=A4,
            rightMargin=15 * mm,
            leftMargin=15 * mm,
            topMargin=15 * mm,
            bottomMargin=15 * mm
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "PetShopTitle",
            parent=styles["Title"],
            fontName=pdf_font,
            fontSize=22,
            leading=26,
            alignment=TA_CENTER,
            spaceAfter=5
        )

        subtitle_style = ParagraphStyle(
            "PetShopSubtitle",
            parent=styles["Normal"],
            fontName=pdf_font,
            fontSize=14,
            leading=18,
            alignment=TA_CENTER,
            spaceAfter=12
        )

        normal_style = ParagraphStyle(
            "PetShopNormal",
            parent=styles["Normal"],
            fontName=pdf_font,
            fontSize=10,
            leading=14
        )

        story = []

        story.append(
            Paragraph(
                "PET SHOP MANAGEMENT SYSTEM",
                title_style
            )
        )

        story.append(
            Paragraph(
                "HÓA ĐƠN BÁN HÀNG",
                subtitle_style
            )
        )

        info_data = [
            [
                Paragraph("Mã hóa đơn:", normal_style),
                Paragraph(invoice["transno"], normal_style),
                Paragraph("Nhân viên:", normal_style),
                Paragraph(invoice["cashier"], normal_style)
            ],
            [
                Paragraph("Khách hàng:", normal_style),
                Paragraph(invoice["customer"], normal_style),
                Paragraph("Trạng thái:", normal_style),
                Paragraph("Đã thanh toán", normal_style)
            ]
        ]

        info_table = Table(
            info_data,
            colWidths=[28 * mm, 62 * mm, 28 * mm, 57 * mm]
        )

        info_table.setStyle(
            TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, 0), (0, -1), colors.lightgrey),
                ("BACKGROUND", (2, 0), (2, -1), colors.lightgrey),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6)
            ])
        )

        story.append(info_table)
        story.append(Spacer(1, 10))

        product_data = [
            [
                Paragraph("STT", normal_style),
                Paragraph("Mã SP", normal_style),
                Paragraph("Tên sản phẩm", normal_style),
                Paragraph("SL", normal_style),
                Paragraph("Đơn giá", normal_style),
                Paragraph("Thành tiền", normal_style)
            ]
        ]

        index = 1

        for item in invoice["items"]:

            product_data.append([
                Paragraph(str(index), normal_style),
                Paragraph(item["pcode"], normal_style),
                Paragraph(item["pname"], normal_style),
                Paragraph(str(item["qty"]), normal_style),
                Paragraph(
                    "{:,.0f}".format(item["price"]),
                    normal_style
                ),
                Paragraph(
                    "{:,.0f}".format(item["total"]),
                    normal_style
                )
            ])

            index = index + 1

        product_table = Table(
            product_data,
            colWidths=[12 * mm, 17 * mm, 65 * mm, 15 * mm, 32 * mm, 38 * mm],
            repeatRows=1
        )

        product_table.setStyle(
            TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("ALIGN", (0, 0), (1, -1), "CENTER"),
                ("ALIGN", (3, 1), (3, -1), "CENTER"),
                ("ALIGN", (4, 1), (5, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6)
            ])
        )

        story.append(product_table)
        story.append(Spacer(1, 10))

        total_data = [
            [
                Paragraph("TỔNG TIỀN", normal_style),
                Paragraph(
                    "{:,.0f} VNĐ".format(invoice["total"]),
                    normal_style
                )
            ]
        ]

        total_table = Table(
            total_data,
            colWidths=[120 * mm, 59 * mm]
        )

        total_table.setStyle(
            TableStyle([
                ("BOX", (0, 0), (-1, -1), 1, colors.grey),
                ("ALIGN", (0, 0), (0, 0), "RIGHT"),
                ("ALIGN", (1, 0), (1, 0), "RIGHT"),
                ("FONTNAME", (0, 0), (-1, -1), pdf_font),
                ("FONTSIZE", (0, 0), (-1, -1), 13),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8)
            ])
        )

        story.append(total_table)
        story.append(Spacer(1, 15))

        qr_image = ReportLabImage(
            qr_file,
            width=38 * mm,
            height=38 * mm
        )

        qr_table = Table(
            [[
                Paragraph(
                    "Quét mã QR để xem thông tin hóa đơn.",
                    normal_style
                ),
                qr_image
            ]],
            colWidths=[120 * mm, 59 * mm]
        )

        qr_table.setStyle(
            TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (1, 0), (1, 0), "CENTER")
            ])
        )

        story.append(qr_table)
        story.append(Spacer(1, 10))
        story.append(
            Paragraph(
                "Cảm ơn quý khách đã mua hàng!",
                subtitle_style
            )
        )

        document.build(story)

        return file_path

    # XUẤT DOCX

    def export_invoice_docx(self, invoice):

        folder = self.get_invoice_folder()

        file_path = os.path.join(
            folder,
            invoice["transno"] + ".docx"
        )

        qr_file = self.create_invoice_qr(invoice)

        document = Document()

        section = document.sections[0]
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = title.add_run(
            "PET SHOP MANAGEMENT SYSTEM"
        )
        run.bold = True
        run.font.size = Pt(20)

        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = subtitle.add_run(
            "HÓA ĐƠN BÁN HÀNG"
        )
        run.bold = True
        run.font.size = Pt(15)

        document.add_paragraph(
            "Mã hóa đơn: "
            + invoice["transno"]
        )

        document.add_paragraph(
            "Khách hàng: "
            + invoice["customer"]
        )

        document.add_paragraph(
            "Nhân viên: "
            + invoice["cashier"]
        )

        table = document.add_table(
            rows=1,
            cols=6
        )

        table.style = "Table Grid"

        headers = [
            "STT",
            "Mã SP",
            "Tên sản phẩm",
            "SL",
            "Đơn giá",
            "Thành tiền"
        ]

        for index in range(6):

            table.rows[0].cells[index].text = headers[index]

        stt = 1

        for item in invoice["items"]:

            row = table.add_row().cells

            row[0].text = str(stt)
            row[1].text = item["pcode"]
            row[2].text = item["pname"]
            row[3].text = str(item["qty"])
            row[4].text = "{:,.0f}".format(item["price"])
            row[5].text = "{:,.0f}".format(item["total"])

            stt = stt + 1

        total_paragraph = document.add_paragraph()
        total_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        run = total_paragraph.add_run(
            "TỔNG TIỀN: {:,.0f} VNĐ".format(
                invoice["total"]
            )
        )
        run.bold = True
        run.font.size = Pt(14)

        qr_paragraph = document.add_paragraph()
        qr_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        qr_paragraph.add_run().add_picture(
            qr_file,
            width=Inches(1.8)
        )

        note = document.add_paragraph()
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = note.add_run(
            "Cảm ơn quý khách đã mua hàng!"
        )
        run.bold = True
        run.font.size = Pt(12)

        document.save(file_path)

        return file_path

    # XỬ LÝ NÚT XUẤT FILE

    def export_invoice_format(
        self,
        transno,
        file_type,
        export_window
    ):

        invoice = self.get_invoice_data(transno)

        if invoice is None:

            messagebox.showerror(
                "Lỗi",
                "Không tìm thấy dữ liệu hóa đơn."
            )

            return

        try:

            file_path = None

            if file_type == "txt":

                file_path = self.export_invoice_txt(
                    invoice
                )

            elif file_type == "pdf":

                file_path = self.export_invoice_pdf(
                    invoice
                )

            elif file_type == "docx":

                file_path = self.export_invoice_docx(
                    invoice
                )

            if file_path is not None:

                messagebox.showinfo(
                    "Xuất hóa đơn thành công",
                    "Đã xuất hóa đơn thành công!\n\n"
                    + "Tệp: "
                    + file_path
                )

                try:
                    export_window.destroy()
                except:
                    pass

        except Exception as e:

            messagebox.showerror(
                "Lỗi xuất hóa đơn",
                "Không thể xuất hóa đơn.\n\n"
                + str(e)
            )
